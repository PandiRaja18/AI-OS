"""Reading a folder of real files: spreadsheets, PDFs, Word documents, text.

Most people's data is not a warehouse, it is a directory. This module turns one
into something the agents can work with:

- tabular files (`.csv`, `.tsv`, `.xlsx`) are ingested into a throwaway SQLite
  database, one table per file or sheet, so the existing query path works
  unchanged;
- document files (`.md`, `.txt`, `.pdf`, `.docx`) are read as text for search.

The ingested database is a copy, which is what makes ad-hoc querying safe to
offer there and nowhere else - nothing the model does can reach the original.
"""

from __future__ import annotations

import csv
import re
import sqlite3
import xml.etree.ElementTree as ElementTree
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

TABULAR = {".csv", ".tsv", ".xlsx", ".xlsm"}
DOCUMENTS = {".md", ".txt", ".pdf", ".docx", ".rst", ".log"}
READABLE = TABULAR | DOCUMENTS

MAX_ROWS_PER_TABLE = 200_000
WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class Table:
    """One ingested table."""

    name: str
    source: str
    columns: list[str]
    rows: int


@dataclass
class Ingestion:
    """The result of turning a folder into a queryable database."""

    database: Path
    tables: list[Table] = field(default_factory=list)
    skipped: list[tuple[str, str]] = field(default_factory=list)

    def describe(self) -> list[dict]:
        return [
            {
                "table": table.name,
                "from_file": table.source,
                "columns": table.columns,
                "rows": table.rows,
            }
            for table in self.tables
        ]


def table_name(stem: str, sheet: str | None = None) -> str:
    """A safe SQL identifier derived from a file name."""
    raw = f"{stem}_{sheet}" if sheet else stem
    cleaned = re.sub(r"[^A-Za-z0-9_]+", "_", raw).strip("_").lower()
    if not cleaned or cleaned[0].isdigit():
        cleaned = f"t_{cleaned}"
    return cleaned[:60]


def column_name(raw: str, index: int) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_]+", "_", str(raw or "")).strip("_").lower()
    if not cleaned or cleaned[0].isdigit():
        cleaned = f"col_{index}" if not cleaned else f"c_{cleaned}"
    return cleaned[:60]


def _unique(name: str, taken: set[str]) -> str:
    candidate, suffix = name, 2
    while candidate in taken:
        candidate = f"{name}_{suffix}"
        suffix += 1
    taken.add(candidate)
    return candidate


def read_csv(path: Path) -> tuple[list[str], list[list]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        rows = [row for _, row in zip(range(MAX_ROWS_PER_TABLE + 1), reader)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def read_xlsx(path: Path) -> dict[str, tuple[list[str], list[list]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, tuple[list[str], list[list]]] = {}
    for sheet in workbook.worksheets:
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index > MAX_ROWS_PER_TABLE:
                break
            rows.append(list(row))
        if rows:
            sheets[sheet.title] = ([str(c) for c in rows[0]], rows[1:])
    workbook.close()
    return sheets


def read_pdf(path: Path) -> str:
    """Text layer only. A scanned PDF needs OCR, which this does not do."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def read_docx(path: Path) -> str:
    """Word documents are a zip of XML, so this needs no extra dependency."""
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    tree = ElementTree.fromstring(xml)
    paragraphs = []
    for node in tree.iter(f"{WORD_NAMESPACE}p"):
        text = "".join(run.text or "" for run in node.iter(f"{WORD_NAMESPACE}t"))
        if text.strip():
            paragraphs.append(text.strip())
    return "\n\n".join(paragraphs)


def read_document(path: Path) -> str:
    """Extract text from any supported document type."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix == ".docx":
        return read_docx(path)
    return path.read_text(encoding="utf-8", errors="replace")


def ingest_folder(folder: Path, database: Path) -> Ingestion:
    """Load every tabular file in `folder` into a fresh SQLite database."""
    database.parent.mkdir(parents=True, exist_ok=True)
    database.unlink(missing_ok=True)
    result = Ingestion(database=database)
    taken: set[str] = set()

    with sqlite3.connect(database) as connection:
        for path in sorted(folder.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in TABULAR:
                continue
            try:
                if path.suffix.lower() in (".xlsx", ".xlsm"):
                    for sheet, (header, rows) in read_xlsx(path).items():
                        _create(
                            connection, result, taken, path, header, rows, sheet
                        )
                else:
                    header, rows = read_csv(path)
                    _create(connection, result, taken, path, header, rows)
            except Exception as error:
                result.skipped.append((path.name, f"{type(error).__name__}: {error}"))
    return result


def _create(
    connection: sqlite3.Connection,
    result: Ingestion,
    taken: set[str],
    path: Path,
    header: list[str],
    rows: list[list],
    sheet: str | None = None,
) -> None:
    if not header:
        result.skipped.append((path.name, "no header row"))
        return
    if any(_unprintable(str(value)) for value in header):
        # A binary file given a .csv name decodes into nonsense rather than
        # failing, so it would otherwise become a junk table nobody can use.
        result.skipped.append((path.name, "header is not text; looks binary"))
        return
    name = _unique(table_name(path.stem, sheet), taken)
    columns = [column_name(value, index) for index, value in enumerate(header)]
    columns = [_unique(column, set()) if columns.count(column) == 1 else f"{column}_{i}"
               for i, column in enumerate(columns)]

    definition = ", ".join(f'"{column}" TEXT' for column in columns)
    connection.execute(f'CREATE TABLE "{name}" ({definition})')
    placeholders = ",".join("?" * len(columns))
    connection.executemany(
        f'INSERT INTO "{name}" VALUES ({placeholders})',
        [_fit(row, len(columns)) for row in rows],
    )
    result.tables.append(
        Table(name=name, source=path.name, columns=columns, rows=len(rows))
    )


def _unprintable(value: str) -> bool:
    """True when text contains control or replacement characters."""
    return any(ord(ch) < 32 and ch not in "	" or ch == "�" for ch in value)


def _fit(row: list, width: int) -> list:
    """Pad or trim a row, and store everything as text for predictable typing."""
    values = [None if value is None else str(value) for value in row[:width]]
    return values + [None] * (width - len(values))


def inventory(folder: Path) -> dict[str, list[str]]:
    """What a folder contains, grouped by how it would be used."""
    tabular, documents, ignored = [], [], []
    for path in sorted(folder.rglob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        target = (
            tabular if suffix in TABULAR
            else documents if suffix in DOCUMENTS
            else ignored
        )
        target.append(path.name)
    return {"tabular": tabular, "documents": documents, "ignored": ignored}
